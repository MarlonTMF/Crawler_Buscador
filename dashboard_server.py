from __future__ import annotations

import json
import threading
import subprocess
import uuid
import time
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
from typing import Dict, Any

from crawler.core.validation_engine import build_record_evidence, build_validation_summary, load_json_records

ROOT = Path(__file__).resolve().parent
DASHBOARD_DIR = ROOT / "dashboard"
JSON_PATH = ROOT / "output" / "excel_urls_diagnostic.json"
BASELINE_PATH = ROOT / "resultadosPrimerCrawleoUnido.xlsx"

# Simple in-memory job registry
JOBS: Dict[str, Dict[str, Any]] = {}
JOBS_LOCK = threading.Lock()
JOBS_DIR = ROOT / 'output' / 'jobs'
JOBS_DIR.mkdir(parents=True, exist_ok=True)


def _read_records_for_summary() -> dict:
    records = load_json_records(JSON_PATH)
    try:
        from benchmark_runner import SOURCES

        existing_urls = {
            record.get('Url_Original') or record.get('url')
            for record in records
            if isinstance(record, dict)
        }
        records.extend({
            'Fuente': source['code'],
            'Institucion': source['name'],
            'Url_Original': source['url'],
            'Score_Excel': 0.0,
            'Diagnosticos_Excel': ['pendiente de análisis'],
            'HTTP_Status': 'PENDING',
            'Final_Url': source['url'],
        } for source in SOURCES if source['url'] not in existing_urls)
    except Exception:
        pass
    baseline_rows = 0
    if BASELINE_PATH.exists():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(BASELINE_PATH, read_only=True, data_only=True)
            ws = wb.active
            baseline_rows = sum(
                1
                for row in ws.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in row)
            )
        except Exception:
            baseline_rows = 0

    summary = build_validation_summary(records, baseline_rows=baseline_rows)

    # load persisted moved-url mappings and annotate records with mapping state
    mappings_path = ROOT / 'config' / 'moved_urls.json'
    mappings = []
    if mappings_path.exists():
        try:
            mappings = json.loads(mappings_path.read_text(encoding='utf-8'))
        except Exception:
            mappings = []

    annotated = []
    for record in records:
        rec = {**record, "evidence": build_record_evidence(record)}
        original = record.get('Url_Original') or record.get('original') or record.get('url')
        matched = None
        for m in mappings:
            if not m:
                continue
            if m.get('original') and original and m.get('original') == original:
                matched = m
        if matched:
            status = 'pending'
            if 'accepted' in matched:
                status = 'accepted' if matched.get('accepted') else 'rejected'
            rec['mapping_status'] = status
            rec['mapping_resolved'] = matched.get('resolved')
            rec['mapping_entry'] = matched
        else:
            rec['mapping_status'] = 'none'
        annotated.append(rec)

    summary["records"] = annotated
    return summary


class DashboardHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(DASHBOARD_DIR), **kwargs)

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/summary":
            payload = json.dumps(_read_records_for_summary()).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)
            return

        # job status / logs
        if parsed.path.startswith('/api/job/'):
            parts = parsed.path.strip('/').split('/')
            # /api/job/{id}
            if len(parts) >= 3 and parts[0] == 'api' and parts[1] == 'job':
                jobid = parts[2]
                with JOBS_LOCK:
                    job = JOBS.get(jobid)
                if not job:
                    progress_path = JOBS_DIR / f'{jobid}.progress.json'
                    if progress_path.exists():
                        try:
                            progress = json.loads(progress_path.read_text(encoding='utf-8'))
                            minimal = {'id': jobid, 'type': 'run_all', 'status': 'running', 'stage': 'processing', 'stage_detail': 'Procesando URLs una por una', **progress}
                            payload = json.dumps(minimal).encode('utf-8')
                            self.send_response(200)
                            self.send_header('Content-Type', 'application/json; charset=utf-8')
                            self.send_header('Content-Length', str(len(payload)))
                            self.end_headers()
                            self.wfile.write(payload)
                            return
                        except Exception:
                            pass
                    log_path = JOBS_DIR / f'{jobid}.log'
                    if log_path.exists():
                        try:
                            alternatives = json.loads(log_path.read_text(encoding='utf-8'))
                            if isinstance(alternatives, list):
                                minimal = {'id': jobid, 'type': 'probe_alternatives', 'status': 'finished', 'stage': 'completed', 'stage_detail': f"{sum(1 for item in alternatives if item.get('reachable'))} alternativas accesibles de {len(alternatives)} probadas", 'alternatives': alternatives}
                                payload = json.dumps(minimal).encode('utf-8')
                                self.send_response(200)
                                self.send_header('Content-Type', 'application/json; charset=utf-8')
                                self.send_header('Content-Length', str(len(payload)))
                                self.end_headers()
                                self.wfile.write(payload)
                                return
                        except Exception:
                            pass
                    meta_path = JOBS_DIR / f'{jobid}.meta.json'
                    if meta_path.exists():
                        try:
                            minimal = json.loads(meta_path.read_text(encoding='utf-8'))
                            payload = json.dumps(minimal).encode('utf-8')
                            self.send_response(200)
                            self.send_header('Content-Type', 'application/json; charset=utf-8')
                            self.send_header('Content-Length', str(len(payload)))
                            self.end_headers()
                            self.wfile.write(payload)
                            return
                        except Exception:
                            pass
                    self.send_response(404)
                    self.end_headers()
                    return
                # /api/job/{id}/logs
                if len(parts) == 4 and parts[3] == 'logs':
                    try:
                        logpath = Path(job.get('log_path'))
                        text = logpath.read_text(encoding='utf-8') if logpath.exists() else ''
                    except Exception:
                        text = ''
                    payload = json.dumps({'job_id': jobid, 'logs': text}).encode('utf-8')
                    self.send_response(200)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.send_header('Content-Length', str(len(payload)))
                    self.end_headers()
                    self.wfile.write(payload)
                    return

                # /api/job/{id}/updates
                if len(parts) == 4 and parts[3] == 'updates':
                    payload = json.dumps({'job_id': jobid, 'updated': job.get('updated_urls', [])}).encode('utf-8')
                    self.send_response(200)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.send_header('Content-Length', str(len(payload)))
                    self.end_headers()
                    self.wfile.write(payload)
                    return

                # default: job metadata
                minimal = {k: job.get(k) for k in ('id', 'type', 'status', 'stage', 'stage_detail', 'alternative_url', 'target_url', 'error_detail', 'alternatives', 'analysis_results', 'tested', 'total', 'start_ts', 'end_ts')}
                progress_path = job.get('progress_path')
                if progress_path:
                    try:
                        minimal.update(json.loads(Path(progress_path).read_text(encoding='utf-8')))
                    except Exception:
                        pass
                if minimal.get('type') == 'run_all' and minimal.get('status') == 'running' and minimal.get('total'):
                    minimal['stage'] = 'processing'
                    minimal['stage_detail'] = f"Analizando URL {minimal.get('processed', 0)} de {minimal['total']}"
                if job.get('error'):
                    minimal['error'] = job['error']
                payload = json.dumps(minimal).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)
                return

        if parsed.path in {"/", "/index.html"}:
            self.path = "/index.html"
        super().do_GET()

    def log_message(self, format, *args):
        # avoid noisy logs in terminal
        return

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path.startswith('/api/mapping/'):
            action = parsed.path.rsplit('/', 1)[-1]
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length) if length else b''
            try:
                payload = json.loads(body.decode('utf-8') or '{}')
            except Exception:
                payload = {}

            original = payload.get('original')
            resolved = payload.get('resolved')

            cfg_path = ROOT / 'config' / 'moved_urls.json'
            data = []
            if cfg_path.exists():
                try:
                    data = json.loads(cfg_path.read_text(encoding='utf-8'))
                except Exception:
                    data = []

            found = False
            for entry in data:
                if entry.get('original') == original and entry.get('resolved') == resolved:
                    entry['accepted'] = True if action == 'accept' else False
                    entry['accepted_at'] = json.dumps({'ts': None})
                    found = True
                    break

            if not found:
                entry = {
                    'original': original,
                    'resolved': resolved,
                    'confidence': 0.8,
                    'timestamp': None,
                    'accepted': True if action == 'accept' else False,
                }
                data.append(entry)

            try:
                cfg_path.parent.mkdir(parents=True, exist_ok=True)
                cfg_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            except Exception:
                pass

            resp = json.dumps({'status': 'ok', 'action': action, 'original': original, 'resolved': resolved}).encode('utf-8')
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(resp)))
            self.end_headers()
            self.wfile.write(resp)
            return

        if parsed.path == '/api/run_url/confirm':
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length) if length else b''
            try:
                payload = json.loads(body.decode('utf-8') or '{}')
            except Exception:
                payload = {}
            jobid = payload.get('job_id')
            with JOBS_LOCK:
                job = JOBS.get(jobid)
                if job and job.get('confirmation_event'):
                    job['confirmed'] = True
                    job['confirmation_event'].set()
            resp = json.dumps({'status': 'accepted', 'job_id': jobid}).encode('utf-8')
            self.send_response(202 if job else 404)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(resp)))
            self.end_headers()
            self.wfile.write(resp)
            return

        if parsed.path == '/api/run_all':
            # start a background job that runs the rebuild script and captures logs
            jobid = uuid.uuid4().hex
            logpath = JOBS_DIR / f"{jobid}.log"

            def _run_all_job(jid, logfile):
                with JOBS_LOCK:
                    JOBS[jid].update({'status': 'running', 'start_ts': time.time()})
                try:
                    with open(logfile, 'w', encoding='utf-8') as fh:
                        env = dict(__import__('os').environ)
                        env['DATAX_PROGRESS_PATH'] = str(JOBS[jid].get('progress_path', ''))
                        proc = subprocess.Popen(["python", str(ROOT / 'scripts' / 'rebuild_excel_diagnostics.py')], cwd=str(ROOT), stdout=fh, stderr=fh, env=env)
                        proc.wait()

                    # compute updated URLs by diffing old/new JSON
                    before = []
                    try:
                        # on job start, store a snapshot path
                        snap = JOBS[jid].get('snapshot_before')
                        if snap and Path(snap).exists():
                            before = json.loads(Path(snap).read_text(encoding='utf-8'))
                    except Exception:
                        before = []

                    after = []
                    try:
                        if JSON_PATH.exists():
                            after = json.loads(JSON_PATH.read_text(encoding='utf-8'))
                    except Exception:
                        after = []

                    before_urls = set((r.get('Url_Original') or r.get('url') or '') for r in before)
                    after_urls = set((r.get('Url_Original') or r.get('url') or '') for r in after)
                    changed = list(after_urls - before_urls)
                    with JOBS_LOCK:
                        JOBS[jid].update({'status': 'finished', 'end_ts': time.time(), 'updated_urls': changed})
                except Exception as e:
                    with JOBS_LOCK:
                        JOBS[jid].update({'status': 'failed', 'end_ts': time.time(), 'error': str(e)})

            # snapshot current JSON for diff
            snap_path = JOBS_DIR / f"{jobid}.before.json"
            try:
                if JSON_PATH.exists():
                    snap_path.write_text(JSON_PATH.read_text(encoding='utf-8'), encoding='utf-8')
            except Exception:
                pass

            with JOBS_LOCK:
                progress_path = JOBS_DIR / f"{jobid}.progress.json"
                progress_path.write_text(json.dumps({'processed': 0, 'total': 0}), encoding='utf-8')
                JOBS[jobid] = {'id': jobid, 'type': 'run_all', 'status': 'queued', 'progress_path': str(progress_path), 'log_path': str(logpath), 'snapshot_before': str(snap_path), 'updated_urls': []}

            t = threading.Thread(target=_run_all_job, args=(jobid, str(logpath)), daemon=True)
            t.start()
            resp = json.dumps({'status': 'accepted', 'action': 'run_all', 'job_id': jobid}).encode('utf-8')
            self.send_response(202)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(resp)))
            self.end_headers()
            self.wfile.write(resp)
            return

        if parsed.path == '/api/run_url':
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length) if length else b''
            try:
                payload = json.loads(body.decode('utf-8') or '{}')
            except Exception:
                payload = {}
            url = payload.get('url')
            source_original = payload.get('original_url') or url
            if url:
                # start a job for single-url validation
                jobid = uuid.uuid4().hex
                logpath = JOBS_DIR / f"{jobid}.log"

                def _run_url_job(jid, logfile, target_url, original_url=None):
                    with JOBS_LOCK:
                        JOBS[jid].update({'status': 'running', 'stage': 'original', 'stage_detail': 'Comprobando la URL original', 'target_url': target_url, 'start_ts': time.time()})
                    try:
                        from crawler.core.fetcher import HttpFetcher

                        fetcher = HttpFetcher(timeout=15, max_retries=2, rate_limit_seconds=0.5)
                        with JOBS_LOCK:
                            JOBS[jid].update({'stage': 'direct_check', 'stage_detail': 'Validando acceso directo, robots.txt y respuesta HTTP'})
                        result = fetcher.validate_url_access(target_url, browser_fallback=True, allow_variants=False)
                        if not result.get('reachable_http'):
                            with JOBS_LOCK:
                                JOBS[jid].update({'stage': 'searching_alternative', 'stage_detail': 'La URL original no respondió. Buscando una alternativa similar.'})
                            alternative_url = fetcher._try_url_variants(target_url)
                            if alternative_url:
                                confirmation_event = threading.Event()
                                with JOBS_LOCK:
                                    JOBS[jid].update({'status': 'awaiting_confirmation', 'stage': 'awaiting_confirmation', 'stage_detail': 'Alternativa sugerida. Esperando tu confirmación para analizarla.', 'alternative_url': alternative_url, 'confirmation_event': confirmation_event})
                                if not confirmation_event.wait(timeout=300):
                                    raise RuntimeError('No se confirmó la URL alternativa dentro del tiempo permitido.')
                                with JOBS_LOCK:
                                    JOBS[jid].update({'status': 'running', 'stage': 'alternative_check', 'stage_detail': f'Analizando alternativa aceptada: {alternative_url}', 'target_url': alternative_url})
                                result = fetcher.validate_url_access(alternative_url, browser_fallback=True, allow_variants=False)
                                result['mapped_from'] = target_url
                                result['resolved_from_variant'] = True
                                result['original_url'] = target_url
                        mapped = bool(result.get('resolved_from_variant'))
                        error_type = result.get('error_type')
                        with JOBS_LOCK:
                            JOBS[jid].update({
                                'stage': 'alternative_resolved' if mapped else ('completed' if not error_type else 'failed'),
                                'stage_detail': (f"Alternativa encontrada y validada: {result.get('final_url')}" if mapped else (f"Error técnico: {error_type}" if error_type else 'Acceso directo validado')),
                                'alternative_url': result.get('final_url') if mapped else None,
                            })
                        with open(logfile, 'w', encoding='utf-8') as fh:
                            fh.write(json.dumps(result, ensure_ascii=False, indent=2))

                        effective_score = float(result.get('effective_score', 0.0) or 0.0)
                        http_status = result.get('status_code')
                        http_status = str(http_status) if http_status not in (None, 0) else (result.get('error_type') or 'CONN_ERROR')
                        reachable = bool(result.get('reachable_http'))
                        has_doc_signal = bool(result.get('has_document_signal'))
                        analyzed_url = result.get('url') or target_url
                        robots_allowed = fetcher.is_url_allowed_by_robots(analyzed_url)
                        doc_links = int(result.get('document_links_found') or 0)
                        keyword_hits = list(result.get('keyword_hits') or [])
                        document_evidence = result.get('document_evidence') or {
                            'keyword_hits': keyword_hits,
                            'file_type': None,
                            'quality_score': effective_score,
                            'snippet_text': result.get('snippet_text') or '',
                        }

                        diagnostics = []
                        if not reachable:
                            diagnostics = ['sin enlace directo', 'fallo operativo']
                        elif has_doc_signal:
                            diagnostics = ['exitoso', 'documento detectado', 'acceso real']
                        else:
                            diagnostics = ['exitoso', 'sin enlace directo']

                        error_detail = result.get('error_type')
                        if error_detail:
                            diagnostics.append(error_detail.lower())

                        # load existing records and update matching entry
                        try:
                            if JSON_PATH.exists():
                                data = json.loads(JSON_PATH.read_text(encoding='utf-8'))
                            else:
                                data = []
                        except Exception:
                            data = []

                        # An individual validation must update one row, never replace the dataset.
                        # Recover the configured source list when a previous run left a partial JSON.
                        try:
                            from benchmark_runner import SOURCES

                            existing_urls = {
                                entry.get('Url_Original') or entry.get('url')
                                for entry in data
                                if isinstance(entry, dict)
                            }
                            if len(data) < len(SOURCES):
                                data.extend({
                                    'Fuente': source['code'],
                                    'Institucion': source['name'],
                                    'Url_Original': source['url'],
                                    'Score_Excel': 0.0,
                                    'Diagnosticos_Excel': ['pendiente de análisis'],
                                    'HTTP_Status': 'PENDING',
                                    'Final_Url': source['url'],
                                } for source in SOURCES if source['url'] not in existing_urls)
                        except Exception:
                            pass

                        found = False
                        record_url = original_url or result.get('original_url') or target_url
                        for entry in data:
                            orig = entry.get('Url_Original') or entry.get('original') or entry.get('url')
                            if orig == record_url:
                                entry.update({
                                    'Score_Excel': round(effective_score, 1),
                                    'Diagnosticos_Excel': diagnostics,
                                    'Robots_Allowed': robots_allowed,
                                    'HTTP_Status': http_status,
                                    'Final_Url': result.get('final_url') or target_url,
                                    'mapped_from': result.get('mapped_from'),
                                    'resolved_from_variant': bool(result.get('resolved_from_variant')),
                                    'original_url': result.get('original_url'),
                                    'mapping_resolved': result.get('final_url') if result.get('resolved_from_variant') else None,
                                    'Doc_Links_Found_In_Seed': doc_links,
                                    'Subpage_Keywords_Found': len(keyword_hits),
                                    'Document_Evidence': document_evidence,
                                    'Error_Detail': error_detail if error_detail else None,
                                })
                                found = True
                                break

                        if not found:
                            new_entry = {
                                'Fuente': 'manual',
                                'Institucion': 'manual',
                                'Url_Original': record_url,
                                'Score_Excel': round(effective_score, 1),
                                'Diagnosticos_Excel': diagnostics,
                                'Robots_Allowed': robots_allowed,
                                'HTTP_Status': http_status,
                                'Final_Url': result.get('final_url') or target_url,
                                'mapped_from': result.get('mapped_from'),
                                'resolved_from_variant': bool(result.get('resolved_from_variant')),
                                'original_url': result.get('original_url'),
                                'mapping_resolved': result.get('final_url') if result.get('resolved_from_variant') else None,
                                'Doc_Links_Found_In_Seed': doc_links,
                                'Subpage_Keywords_Found': len(keyword_hits),
                                'Document_Evidence': document_evidence,
                                'Error_Detail': error_detail if error_detail else None,
                            }
                            data.append(new_entry)

                        try:
                            JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
                            JSON_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
                        except Exception:
                            pass

                        # record updated URL for reporting
                        with JOBS_LOCK:
                            JOBS[jid].update({'status': 'finished', 'stage': 'completed', 'stage_detail': (f"Alternativa validada: {result.get('final_url')}" if mapped else (f"No se pudo conectar: {error_type}" if error_type else 'URL validada')), 'end_ts': time.time(), 'updated_urls': [target_url], 'error_detail': error_type})
                    except Exception as e:
                        with JOBS_LOCK:
                            JOBS[jid].update({'status': 'failed', 'stage': 'failed', 'stage_detail': f'Error durante el análisis: {e}', 'end_ts': time.time(), 'error': str(e), 'error_detail': str(e)})

                with JOBS_LOCK:
                    JOBS[jobid] = {'id': jobid, 'type': 'run_url', 'status': 'queued', 'log_path': str(logpath), 'updated_urls': []}

                t = threading.Thread(target=_run_url_job, args=(jobid, str(logpath), url, source_original), daemon=True)
                t.start()
                resp = json.dumps({'status': 'accepted', 'action': 'run_url', 'job_id': jobid, 'url': url}).encode('utf-8')
                self.send_response(202)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(resp)))
                self.end_headers()
                self.wfile.write(resp)
                return

        if parsed.path == '/api/probe_alternatives':
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length) if length else b''
            try:
                payload = json.loads(body.decode('utf-8') or '{}')
            except Exception:
                payload = {}
            url = payload.get('url')
            if url:
                jobid = uuid.uuid4().hex
                logpath = JOBS_DIR / f"{jobid}.log"
                metapath = JOBS_DIR / f"{jobid}.meta.json"
                initial_job = {'id': jobid, 'type': 'probe_alternatives', 'status': 'queued', 'stage': 'queued', 'stage_detail': 'Preparando variantes', 'target_url': url, 'alternatives': []}
                metapath.write_text(json.dumps(initial_job, ensure_ascii=False), encoding='utf-8')
                with JOBS_LOCK:
                    JOBS[jobid] = {**initial_job, 'meta_path': str(metapath), 'log_path': str(logpath), 'tested': 0, 'total': 0}

                def _probe_job(jid, logfile, target_url):
                    with JOBS_LOCK:
                        JOBS[jid].update({'status': 'running', 'stage': 'probing', 'stage_detail': 'Probando esquema, www, ruta y terminaciones .org/.bo/.com'})
                        metapath = Path(JOBS[jid]['meta_path'])
                        metapath.write_text(json.dumps(JOBS[jid], ensure_ascii=False), encoding='utf-8')
                    try:
                        from crawler.core.fetcher import HttpFetcher

                        fetcher = HttpFetcher(timeout=4, max_retries=1, rate_limit_seconds=0.0)
                        candidates = fetcher._generate_url_variants(target_url)
                        with JOBS_LOCK:
                            JOBS[jid].update({'total': len(candidates)})
                        results = []
                        from concurrent.futures import ThreadPoolExecutor, as_completed

                        def probe(candidate):
                            local_fetcher = HttpFetcher(timeout=4, max_retries=1, rate_limit_seconds=0.0)
                            return local_fetcher.probe_url_variant(candidate)

                        with ThreadPoolExecutor(max_workers=6) as executor:
                            futures = {executor.submit(probe, candidate): candidate for candidate in candidates}
                            for index, future in enumerate(as_completed(futures), start=1):
                                item = future.result()
                                results.append(item)
                                with JOBS_LOCK:
                                    JOBS[jid].update({'tested': index, 'alternatives': list(results), 'stage_detail': f'Probada {index} de {len(candidates)}: {item["url"]}'})
                                    Path(JOBS[jid]['meta_path']).write_text(json.dumps(JOBS[jid], ensure_ascii=False), encoding='utf-8')
                        with open(logfile, 'w', encoding='utf-8') as fh:
                            fh.write(json.dumps(results, ensure_ascii=False, indent=2))
                        reachable = [item for item in results if item.get('reachable')]
                        with JOBS_LOCK:
                            JOBS[jid].update({'status': 'finished', 'stage': 'completed', 'stage_detail': f'{len(reachable)} alternativas accesibles de {len(results)} probadas', 'alternatives': results, 'end_ts': time.time()})
                            Path(JOBS[jid]['meta_path']).write_text(json.dumps(JOBS[jid], ensure_ascii=False), encoding='utf-8')
                    except Exception as exc:
                        with JOBS_LOCK:
                            JOBS[jid].update({'status': 'failed', 'stage': 'failed', 'stage_detail': f'Error probando alternativas: {exc}', 'error_detail': str(exc), 'end_ts': time.time()})
                            Path(JOBS[jid]['meta_path']).write_text(json.dumps(JOBS[jid], ensure_ascii=False), encoding='utf-8')

                threading.Thread(target=_probe_job, args=(jobid, str(logpath), url), daemon=True).start()
                resp = json.dumps({'status': 'accepted', 'action': 'probe_alternatives', 'job_id': jobid, 'url': url}).encode('utf-8')
                self.send_response(202)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(resp)))
                self.end_headers()
                self.wfile.write(resp)
                return

        if parsed.path == '/api/analyze_alternatives':
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length) if length else b''
            try:
                payload = json.loads(body.decode('utf-8') or '{}')
            except Exception:
                payload = {}
            original_url = payload.get('original_url')
            alternatives = payload.get('alternatives') or []
            if original_url and alternatives:
                jobid = uuid.uuid4().hex
                logpath = JOBS_DIR / f'{jobid}.log'
                metapath = JOBS_DIR / f'{jobid}.meta.json'
                initial_job = {'id': jobid, 'type': 'analyze_alternatives', 'status': 'queued', 'stage': 'queued', 'stage_detail': 'Preparando análisis documental', 'target_url': original_url, 'analysis_results': [], 'tested': 0, 'total': len(alternatives)}
                metapath.write_text(json.dumps(initial_job, ensure_ascii=False), encoding='utf-8')
                with JOBS_LOCK:
                    JOBS[jobid] = {**initial_job, 'meta_path': str(metapath), 'log_path': str(logpath)}

                def _analyze_job(jid, logfile, source_url, items):
                    try:
                        from crawler.core.fetcher import HttpFetcher
                        results = []
                        with JOBS_LOCK:
                            JOBS[jid].update({'status': 'running', 'stage': 'document_analysis', 'stage_detail': 'Analizando documentos, palabras clave y señales de relevancia'})
                        for index, item in enumerate(items, start=1):
                            fetcher = HttpFetcher(timeout=8, max_retries=1, rate_limit_seconds=0.0)
                            detail = fetcher.validate_url_access(item['url'], browser_fallback=True, allow_variants=False)
                            result = {**item, 'score': float(detail.get('effective_score', 0) or 0), 'document_links': int(detail.get('document_links_found', 0) or 0), 'keyword_hits': list(detail.get('keyword_hits') or []), 'has_document_signal': bool(detail.get('has_document_signal')), 'http_status': detail.get('status_code'), 'error': detail.get('error_type')}
                            result['relevance'] = result['document_links'] * 4 + len(result['keyword_hits']) * 2 + result['score']
                            results.append(result)
                            results.sort(key=lambda entry: entry['relevance'], reverse=True)
                            with JOBS_LOCK:
                                JOBS[jid].update({'tested': index, 'analysis_results': list(results), 'stage_detail': f'Analizada {index} de {len(items)}: {item["url"]}'})
                                Path(JOBS[jid]['meta_path']).write_text(json.dumps(JOBS[jid], ensure_ascii=False), encoding='utf-8')
                        with open(logfile, 'w', encoding='utf-8') as fh:
                            fh.write(json.dumps(results, ensure_ascii=False, indent=2))
                        with JOBS_LOCK:
                            JOBS[jid].update({'status': 'finished', 'stage': 'completed', 'stage_detail': f'{len(results)} alternativas ordenadas por relevancia', 'analysis_results': results, 'end_ts': time.time()})
                            Path(JOBS[jid]['meta_path']).write_text(json.dumps(JOBS[jid], ensure_ascii=False), encoding='utf-8')
                    except Exception as exc:
                        with JOBS_LOCK:
                            JOBS[jid].update({'status': 'failed', 'stage': 'failed', 'stage_detail': f'Error analizando alternativas: {exc}', 'error_detail': str(exc), 'end_ts': time.time()})
                            Path(JOBS[jid]['meta_path']).write_text(json.dumps(JOBS[jid], ensure_ascii=False), encoding='utf-8')

                threading.Thread(target=_analyze_job, args=(jobid, str(logpath), original_url, alternatives), daemon=True).start()
                resp = json.dumps({'status': 'accepted', 'action': 'analyze_alternatives', 'job_id': jobid}).encode('utf-8')
                self.send_response(202)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(resp)))
                self.end_headers()
                self.wfile.write(resp)
                return

        if parsed.path == '/api/apply_alternative':
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length) if length else b''
            payload = json.loads(body.decode('utf-8') or '{}')
            original = payload.get('original_url')
            selected = payload.get('selected') or {}
            resolved = selected.get('url')
            data = json.loads(JSON_PATH.read_text(encoding='utf-8')) if JSON_PATH.exists() else []
            for entry in data:
                if (entry.get('Url_Original') or entry.get('url')) == original:
                    entry.update({'Url_Original': resolved, 'Final_Url': resolved, 'original_url': original, 'mapped_from': original, 'resolved_from_variant': True, 'mapping_resolved': resolved, 'HTTP_Status': selected.get('http_status') or selected.get('status') or 200, 'Score_Excel': round(float(selected.get('score', 0) or 0), 1), 'Error_Detail': selected.get('error'), 'Diagnosticos_Excel': ['alternativa seleccionada', 'acceso validado']})
                    break
            JSON_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            resp = json.dumps({'status': 'ok', 'original_url': original, 'resolved_url': resolved}).encode('utf-8')
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(resp)))
            self.end_headers()
            self.wfile.write(resp)
            return

        # fallback: unhandled POST
        self.send_response(404)
        self.end_headers()


if __name__ == "__main__":
    server = ThreadingHTTPServer(("127.0.0.1", 8000), DashboardHandler)
    print("Dashboard activo en http://127.0.0.1:8000")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("Apagando dashboard...")
    finally:
        server.server_close()
