from __future__ import annotations

import json
import os
import tkinter as tk
from pathlib import Path
from tkinter import ttk, filedialog, messagebox

from crawler.core.validation_engine import build_validation_summary, load_json_records


class ValidationWindow(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Crawler Validation Dashboard")
        self.geometry("980x640")
        self.minsize(900, 560)

        self.base_dir = Path(__file__).resolve().parents[1]
        self.default_json = self.base_dir / "output" / "excel_urls_diagnostic.json"
        self.default_baseline = self.base_dir / "resultadosPrimerCrawleoUnido.xlsx"

        self._build_ui()

    def _build_ui(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)

        top = ttk.Frame(self, padding=12)
        top.grid(row=0, column=0, sticky="ew")
        top.columnconfigure(1, weight=1)

        ttk.Label(top, text="Archivo de validación JSON:", font=("Segoe UI", 10, "bold")).grid(row=0, column=0, sticky="w", padx=(0, 8))
        self.json_path_var = tk.StringVar(value=str(self.default_json))
        ttk.Entry(top, textvariable=self.json_path_var).grid(row=0, column=1, sticky="ew")
        ttk.Button(top, text="Buscar", command=self.select_json).grid(row=0, column=2, padx=(8, 0))

        ttk.Label(top, text="Archivo base (opcional Excel/JSON):", font=("Segoe UI", 10, "bold")).grid(row=1, column=0, sticky="w", padx=(0, 8), pady=(8, 0))
        self.baseline_path_var = tk.StringVar(value=str(self.default_baseline))
        ttk.Entry(top, textvariable=self.baseline_path_var).grid(row=1, column=1, sticky="ew", pady=(8, 0))
        ttk.Button(top, text="Base", command=self.select_baseline).grid(row=1, column=2, padx=(8, 0), pady=(8, 0))

        ttk.Button(top, text="Ejecutar validación", command=self.run_validation).grid(row=2, column=1, sticky="w", pady=(12, 0))

        self.tree = ttk.Treeview(self, columns=("métrica", "valor"), show="headings")
        self.tree.heading("métrica", text="Métrica")
        self.tree.heading("valor", text="Valor")
        self.tree.column("métrica", width=420, anchor="w")
        self.tree.column("valor", width=220, anchor="e")
        self.tree.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 12))

        self.text = tk.Text(self, wrap="word", state="disabled", bg="#f7f7f7")
        self.text.grid(row=2, column=0, sticky="nsew", padx=12, pady=(0, 12))

        self._populate_default_summary()

    def select_json(self):
        path = filedialog.askopenfilename(filetypes=[("JSON files", "*.json"), ("All files", "*.*")])
        if path:
            self.json_path_var.set(path)

    def select_baseline(self):
        path = filedialog.askopenfilename(filetypes=[("JSON files", "*.json"), ("Excel files", "*.xlsx"), ("All files", "*.*")])
        if path:
            self.baseline_path_var.set(path)

    def _populate_default_summary(self):
        try:
            self.run_validation()
        except Exception:
            pass

    def run_validation(self):
        json_path = self.json_path_var.get()
        baseline_path = self.baseline_path_var.get().strip()

        if not json_path or not os.path.exists(json_path):
            messagebox.showerror("Archivo requerido", "Selecciona un JSON válido para validar.")
            return

        records = load_json_records(json_path)
        baseline_rows = 0
        if baseline_path and os.path.exists(baseline_path):
            if baseline_path.lower().endswith(".xlsx"):
                try:
                    from openpyxl import load_workbook
                    wb = load_workbook(baseline_path, read_only=True, data_only=True)
                    ws = wb.active
                    baseline_rows = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in _))
                except Exception:
                    baseline_rows = 0
            else:
                baseline_rows = len(load_json_records(baseline_path))

        summary = build_validation_summary(records, baseline_rows=baseline_rows)
        self._render_summary(summary)

    def _render_summary(self, summary: dict):
        for item in self.tree.get_children():
            self.tree.delete(item)

        rows = [
            ("Total de registros", summary["total_records"]),
            ("Casos exitosos", summary["positive_cases"]),
            ("Tasa de éxito", f"{summary['success_rate']:.2%}"),
            ("Estado", summary["status"]),
            ("Delta vs baseline", summary["baseline_delta"]),
        ]
        for metric, value in rows:
            self.tree.insert("", tk.END, values=(metric, value))

        details = []
        details.append("=== RESUMEN DE VALIDACIÓN ===")
        details.append(f"Total de registros: {summary['total_records']}")
        details.append(f"Casos exitosos: {summary['positive_cases']}")
        details.append(f"Tasa de éxito: {summary['success_rate']:.2%}")
        details.append(f"Estado: {summary['status']}")
        details.append(f"Delta vs baseline: {summary['baseline_delta']}")
        details.append("")
        details.append("Distribución por score:")
        for score, count in summary["score_distribution"].items():
            details.append(f"  - Score {score}: {count} registros")
        details.append("")
        details.append("Debilidades:")
        if summary["weaknesses"]:
            for item in summary["weaknesses"]:
                details.append(f"  - {item}")
        else:
            details.append("  - Ninguna debilidad dominante detectada")

        self.text.configure(state="normal")
        self.text.delete("1.0", tk.END)
        self.text.insert("1.0", "\n".join(details))
        self.text.configure(state="disabled")


def main():
    app = ValidationWindow()
    app.mainloop()


if __name__ == "__main__":
    main()
